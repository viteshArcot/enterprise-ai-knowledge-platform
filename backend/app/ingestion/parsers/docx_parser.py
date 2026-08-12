"""DOCX text extraction backed by python-docx."""

from io import BytesIO

from docx import Document as DocxDocument

from app.core.exceptions import IngestionError
from app.ingestion.parsers.base import DocumentParser, ParsedDocument


class DocxParser(DocumentParser):
    """Extract paragraph text and core properties from Word documents."""

    extensions = ("docx",)

    def parse(self, content: bytes, file_name: str) -> ParsedDocument:
        """Read a DOCX archive directly from memory."""
        from app.ingestion.parsers.base import ParsedPage
        try:
            document = DocxDocument(BytesIO(content))
        except Exception as exc:
            raise IngestionError(f"Unable to extract text from '{file_name}'.") from exc
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        full_text = "\n\n".join(paragraphs)
        return ParsedDocument(
            text=full_text,
            pages=[ParsedPage(page_number=1, text=full_text)],
            metadata={"author": document.core_properties.author or None},
        )
